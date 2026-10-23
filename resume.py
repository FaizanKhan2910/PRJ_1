
import os
import json
import time
from pathlib import Path
from typing import Optional

from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# -------------------- SETUP --------------------

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY .env file mein nahi mili!")

client = Groq(api_key=my_api_key)
model = "openai/gpt-oss-20b"


# -------------------- JOB DESCRIPTION --------------------

job_descriptions = """
Job Title: Software Engineer I
Company: PayPal
Location: Bengaluru, India (Illustrative)
Experience: 0-2 years (Illustrative)

Education:
Bachelor's degree in Computer Science, Information Technology,
Software Engineering, or a related field.

Responsibilities:
1. Design, develop, test, debug, and maintain software applications.
2. Write clean, reusable code using Java, Python, C++, or similar languages.
3. Develop and integrate RESTful APIs.
4. Implement business logic and data-processing functionality.
5. Work with relational and NoSQL databases.
6. Apply object-oriented programming and software design patterns.
"""


# -------------------- PYDANTIC MODELS --------------------

class JobD(BaseModel):
    role: str
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: Optional[float] = None
    education_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


class Experience(BaseModel):
    company: Optional[str] = None
    role: Optional[str] = None
    duration: Optional[str] = None
    description: Optional[str] = None
    skills_used: list[str] = Field(default_factory=list)


class Education(BaseModel):
    institution: Optional[str] = None
    degree: Optional[str] = None
    field_of_study: Optional[str] = None
    graduation_year: Optional[int] = None
    grade: Optional[str] = None


class Project(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    technologies: list[str] = Field(default_factory=list)
    url: Optional[str] = None


class Certification(BaseModel):
    name: Optional[str] = None
    issuer: Optional[str] = None
    year: Optional[int] = None


class Resume(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    location: Optional[str] = None
    linkedin: Optional[str] = None
    github: Optional[str] = None
    portfolio: Optional[str] = None

    professional_summary: Optional[str] = None
    current_role: Optional[str] = None
    total_experience_years: Optional[float] = None

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    certifications: list[Certification] = Field(default_factory=list)
    achievements: list[str] = Field(default_factory=list)
    languages: list[str] = Field(default_factory=list)


class MatchResult(BaseModel):
    score: float = Field(ge=0, le=100)
    details: dict = Field(default_factory=dict)


job_schema = JobD.model_json_schema()
resume_schema = Resume.model_json_schema()
match_schema = MatchResult.model_json_schema()


# -------------------- GROQ JSON HELPER --------------------

def call_json_model(messages):
    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"},
        temperature=0
    )

    raw_output = response.choices[0].message.content

    if not raw_output:
        raise ValueError("AI ne empty response return kiya.")

    try:
        return json.loads(raw_output)
    except json.JSONDecodeError as e:
        print("Invalid JSON returned by AI:", repr(raw_output))
        raise e


# -------------------- PARSE JOB DESCRIPTION --------------------

def parse_job_description(description):
    messages = [
        {
            "role": "system",
            "content": f"""
            Extract job information and return only JSON matching
            this schema:
            {json.dumps(job_schema)}

            Rules:
            - Do not invent information.
            - Use null for unknown minimum experience.
            - Use empty lists for missing list fields.
            - Separate required and preferred skills.
            - Extract the minimum experience from the description.
            """
        },
        {
            "role": "user",
            "content": description
        }
    ]

    data = call_json_model(messages)
    return JobD(**data)


# -------------------- RESUME TEXT EXTRACTION --------------------

def read_pdf(file_path):
    reader = PdfReader(str(file_path))
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def read_docx(file_path):
    document = Document(str(file_path))
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return read_pdf(file_path)

    if suffix == ".docx":
        return read_docx(file_path)

    return None


# -------------------- PARSE RESUME --------------------

def parse_resume(resume_text):
    if not resume_text or not resume_text.strip():
        raise ValueError("Resume se text extract nahi hua.")

    messages = [
        {
            "role": "system",
            "content": f"""
            You are an expert resume parser.

            Extract structured information from the resume.
            Return only JSON matching this schema:
            {json.dumps(resume_schema)}

            Rules:
            - Do not invent information.
            - Use null for missing optional fields.
            - Use empty lists for missing list fields.
            - Include internships in experiences.
            - Extract skills from the entire resume.
            - Keep projects separate from work experience.
            - Do not assume a graduation year or experience duration.
            - Return a JSON object, not a JSON string.
            """
        },
        {
            "role": "user",
            "content": f"Parse this resume:\n\n{resume_text}"
        }
    ]

    data = call_json_model(messages)
    resume = Resume(**data)

    # Important: return the actual Resume object.
    return resume


# -------------------- MATCH RESUME WITH JOB --------------------

def final_score(job, resume):
    messages = [
        {
            "role": "system",
            "content": f"""
            You are a careful technical recruiter.

            Compare the candidate resume with the job description.

            Return JSON matching this schema:
            {json.dumps(match_schema)}

            The top-level object must contain:
            - score: a number from 0 to 100
            - details: an object containing candidate_name,
              matching_skills, missing_skills,
              experience_requirement_met, education_match,
              project_relevance, and verdict.

            Scoring guidelines:
            - Compare actual skills against job requirements.
            - Consider relevant work experience and projects.
            - Consider education requirements.
            - Do not award skills that are not present in the resume.
            - Do not penalize a candidate merely because a skill is
              mentioned using a different but equivalent phrase.
            - Do not invent candidate information.
            - Treat missing information as unknown, not as proof
              that the candidate lacks the qualification.
            - Keep the verdict concise.
            """
        },
        {
            "role": "user",
            "content": f"""
            JOB DESCRIPTION:
            {job.model_dump_json(indent=2)}

            CANDIDATE RESUME:
            {resume.model_dump_json(indent=2)}
            """
        }
    ]

    data = call_json_model(messages)
    return MatchResult(**data)


# -------------------- MAIN PROGRAM --------------------

def main():
    resume_folder = Path("resumes")

    if not resume_folder.exists():
        print("Resume folder nahi mila:", resume_folder.resolve())
        print("Project folder ke andar 'resumes' folder create karo.")
        return

    # Parse JD once, then reuse it for all candidates.
    print("Parsing job description...")
    job = parse_job_description(job_descriptions)

    print("Role:", job.role)
    print("Required skills:", job.required_skills)

    all_results = []

    resume_files = sorted(
        file_path for file_path in resume_folder.iterdir()
        if file_path.is_file()
        and file_path.suffix.lower() in [".pdf", ".docx"]
    )

    if not resume_files:
        print("Resumes folder mein koi PDF ya DOCX file nahi mili.")
        return

    for file_path in resume_files:
        print(f"\nProcessing: {file_path.name}")

        try:
            resume_text = read_resume(file_path)

            if not resume_text or not resume_text.strip():
                print("Skipped: file se readable text nahi mila.")
                continue

            parsed_resume = parse_resume(resume_text)

            print("Candidate:", parsed_resume.name or file_path.stem)

            # Keep a short pause between API requests.
            time.sleep(2)

            result = final_score(job, parsed_resume)

            print("Score:", result.score, "%")
            print("Details:", result.details)

            all_results.append({
                "name": parsed_resume.name or file_path.stem,
                "score": result.score,
                "details": result.details
            })

            time.sleep(2)

        except Exception as e:
            print(f"Failed to process {file_path.name}: {e}")
            continue

    if not all_results:
        print("\nKisi bhi resume ko successfully process nahi kiya gaya.")
        return

    all_results.sort(
        key=lambda candidate: candidate["score"],
        reverse=True
    )

    print("\n========== TOP 2 CANDIDATES ==========")

    for candidate in all_results[:2]:
        print(
            candidate["name"],
            "-",
            candidate["score"],
            "%"
        )
        print(candidate["details"])

    print("\n========== LOWEST 2 CANDIDATES ==========")

    for candidate in reversed(all_results[-2:]):
        print(
            candidate["name"],
            "-",
            candidate["score"],
            "%"
        )
        print(candidate["details"])


if __name__ == "__main__":
    main()
